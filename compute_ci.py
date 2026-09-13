"""Compute the Ci parameter for files 16, 17, 18 and add it to their results.

Run `run.py` first — this adds to the results/<N> - raw_result.xlsx
workbooks it already produces, rather than regenerating them.

Ci combines each cycle's Strain (from the Time/Stress/Strain file) with
the matching Extension reading from the paired LVDT file
(Time/Extension), at the moment Stress crosses from positive to negative
within that cycle:

    Ci = (LVDT extension / 2) - (strain_i * 70)

For each cycle already detected in the Stress signal (Max at max_time,
Min at min_time, with Max > 0 and Min < 0 — cycles that don't straddle
zero are skipped, since there's no descending zero-crossing to find):

1. Within the descending window [max_time, min_time], find the raw
   sample where Stress is closest to zero while still positive (the last
   point before it goes negative on the way down).
2. Read Strain at that same sample (strain_i) — same row, same file.
3. In the paired LVDT file (its own, independently-sampled Time column),
   find the closest Time to that same point and read Extension there.
4. Ci = (Extension / 2) - (strain_i * 70)

Adds a "Ci" sheet (Cycle, Time, Strain, LVDT Extension, Ci) and a
"Ci vs Time" chart to the existing output workbook.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from cyclic_loading_analyzer.cli import cycles_for_signal
from cyclic_loading_analyzer.detector import Cycle
from cyclic_loading_analyzer.io_utils import read_raw_data, read_single_signal_data

FILE_NUMBERS = [16, 17, 18]
SHOW_LAST_CYCLE = True
STRESS_THRESHOLD = None  # None = auto-compute, matching run.py

PROJECT_DIR = Path(__file__).parent
RAW_DIR = PROJECT_DIR / "raw_data"
RESULTS_DIR = PROJECT_DIR / "results"

HEADER_FONT = Font(bold=True)
CI_SERIES_COLOR = "1F77B4"


def _find_raw_and_lvdt_paths(number: int) -> tuple[Path, Path]:
    prefix = str(number)
    raw_path = lvdt_path = None
    for p in RAW_DIR.iterdir():
        if not p.name.startswith(prefix):
            continue
        rest = p.name[len(prefix) :]
        if rest and rest[0].isdigit():
            continue  # e.g. "160" shouldn't match number=16
        if "lvdt" in p.name.lower():
            lvdt_path = p
        else:
            raw_path = p
    if raw_path is None or lvdt_path is None:
        raise FileNotFoundError(
            f"Could not find both a raw and an lvdt file for {number!r} in {RAW_DIR}"
        )
    return raw_path, lvdt_path


def _zero_crossing_strain(
    time: np.ndarray, stress: np.ndarray, strain: np.ndarray, max_time: float, min_time: float
) -> tuple[float, float] | None:
    """The last point before Stress crosses from positive to negative
    within [max_time, min_time] — closest to zero while still positive."""
    mask = (time >= max_time) & (time <= min_time)
    idxs = np.where(mask)[0]
    positive = idxs[stress[idxs] > 0]
    if len(positive) == 0:
        return None
    best = positive[np.argmin(stress[positive])]
    return float(time[best]), float(strain[best])


def _nearest_value(target_time: float, other_time: np.ndarray, other_value: np.ndarray) -> float:
    idx = np.searchsorted(other_time, target_time)
    candidates = [i for i in (idx - 1, idx) if 0 <= i < len(other_time)]
    best = min(candidates, key=lambda i: abs(other_time[i] - target_time))
    return float(other_value[best])


def compute_ci_rows(raw_path: Path, lvdt_path: Path) -> list[tuple[int, float, float, float, float]]:
    raw_df = read_raw_data(raw_path)
    time = raw_df["Time"].to_numpy()
    stress = raw_df["Stress"].to_numpy()
    strain = raw_df["Strain"].to_numpy()

    cycles: list[Cycle]
    cycles, _threshold, _merged = cycles_for_signal(time, stress, STRESS_THRESHOLD, SHOW_LAST_CYCLE)

    lvdt_df, signal_name = read_single_signal_data(lvdt_path)
    lvdt_time = lvdt_df["Time"].to_numpy()
    lvdt_value = lvdt_df[signal_name].to_numpy()

    rows = []
    for cycle in cycles:
        if cycle.max_value <= 0 or cycle.min_value >= 0:
            continue
        crossing = _zero_crossing_strain(time, stress, strain, cycle.max_time, cycle.min_time)
        if crossing is None:
            continue
        t_zero, strain_i = crossing
        lvdt_ext = _nearest_value(t_zero, lvdt_time, lvdt_value)
        ci = (lvdt_ext / 2) - (strain_i * 70)
        rows.append((cycle.cycle_number, t_zero, strain_i, lvdt_ext, ci))
    return rows


def _write_ci_sheet_and_chart(output_path: Path, rows: list[tuple[int, float, float, float, float]]) -> None:
    wb = load_workbook(output_path)
    if "Ci" in wb.sheetnames:
        del wb["Ci"]
    ws = wb.create_sheet("Ci")

    headers = ["Cycle", "Time", "Strain", "LVDT Extension", "Ci"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(list(row))
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    ws.freeze_panes = "A2"

    charts_ws = wb["Charts"]
    next_row = 1 + 25 * len(charts_ws._charts)

    chart = ScatterChart()
    chart.title = "Ci vs Time"
    chart.style = 13
    chart.x_axis.title = "Time"
    chart.y_axis.title = "Ci"
    chart.width = 24
    chart.height = 12

    last_row = 1 + len(rows)
    xvalues = Reference(ws, min_col=2, min_row=2, max_row=last_row)
    yvalues = Reference(ws, min_col=5, min_row=1, max_row=last_row)
    series = Series(yvalues, xvalues, title_from_data=True)
    series.marker.symbol = "circle"
    series.marker.size = 6
    series.marker.graphicalProperties.solidFill = CI_SERIES_COLOR
    series.marker.graphicalProperties.line.solidFill = CI_SERIES_COLOR
    series.smooth = False
    series.graphicalProperties.line.width = 15000
    series.graphicalProperties.line.solidFill = CI_SERIES_COLOR
    chart.series.append(series)
    chart.legend.position = "b"
    chart.legend.overlay = False

    charts_ws.add_chart(chart, f"A{next_row}")
    wb.save(output_path)


def main() -> int:
    for number in FILE_NUMBERS:
        raw_path, lvdt_path = _find_raw_and_lvdt_paths(number)
        output_path = RESULTS_DIR / f"{raw_path.stem}_result.xlsx"
        if not output_path.exists():
            raise SystemExit(f"{output_path} not found — run run.py first.")

        rows = compute_ci_rows(raw_path, lvdt_path)
        _write_ci_sheet_and_chart(output_path, rows)
        print(f"{raw_path.name}: added Ci sheet with {len(rows)} rows -> {output_path.name}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
