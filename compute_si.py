"""Compute the Si parameter and add it to each file's results.

Run `run.py` first — this adds to the results/<N> - raw_result.xlsx
workbooks it already produces, rather than regenerating them.

Two different formulas are used, depending on the file:

Files 16, 17, 18 — Si anchored on the Stress signal's descending zero
crossing:

    Si = (Extension / 2) - (strain_i * 70)

For each cycle already detected in the Stress signal (Max at max_time,
Min at min_time, with Max > 0 and Min < 0 — cycles that don't straddle
zero are skipped, since there's no descending zero-crossing to find):

1. Within the descending window [max_time, min_time], find the raw
   sample where Stress is closest to zero while still positive (the last
   point before it goes negative on the way down).
2. Read Strain at that same sample (strain_i) — same row, same file.
3. In the Extension signal's own (independently-sampled) Time column,
   find the closest Time to that same point and read Extension there.
4. Si = (Extension / 2) - (strain_i * 70)

Files 19 and up — Si anchored on the Extension (LVDT) signal's own
Max per cycle instead:

    Si = (LVDT_max / 2) - (strain_at_LVDT_max_time * 70)

For each cycle detected in the Extension signal itself (its own
peak-to-peak Max, not a Stress zero-crossing):

1. Take that cycle's own Max value and Max time.
2. In the main Time column, find the closest Time to the LVDT Max time
   and read Strain there.
3. Si = (LVDT_max / 2) - (strain_at_that_time * 70)

Files 19-22 use paired cycle labeling ("1_1", "1_2", "2_1", ...) for the
"Cycle" column, matching the convention already used for those files'
other sheets (see `cross_reference.paired_cycle_label`) — every two
consecutive detected half-cycles are one physical loading cycle.

Adds an "Si" sheet (Cycle, Time, Strain, Extension, Si) and an
"Si vs Time" chart to the existing output workbook.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from cyclic_loading_analyzer.cli import cycles_for_signal
from cyclic_loading_analyzer.cross_reference import paired_cycle_label
from cyclic_loading_analyzer.detector import Cycle
from cyclic_loading_analyzer.io_utils import read_combined_data

ZERO_CROSSING_FILE_NUMBERS = {16, 17, 18}
LVDT_MAX_FILE_NUMBERS = {19, 20, 21, 22, 23, 24, 25, 26, 27, 28}
PAIRED_CYCLE_FILE_NUMBERS = {19, 20, 21, 22}
SHOW_LAST_CYCLE = True
STRESS_THRESHOLD = None  # None = auto-compute, matching run.py
EXTENSION_THRESHOLD = None  # None = auto-compute, matching run.py

PROJECT_DIR = Path(__file__).parent
RAW_DIR = PROJECT_DIR / "raw_data"
RESULTS_DIR = PROJECT_DIR / "results"

HEADER_FONT = Font(bold=True)
SI_SERIES_COLOR = "1F77B4"


def _find_raw_path(number: int) -> Path:
    prefix = str(number)
    for p in RAW_DIR.iterdir():
        if not p.name.startswith(prefix):
            continue
        rest = p.name[len(prefix) :]
        if rest and rest[0].isdigit():
            continue  # e.g. "160" shouldn't match number=16
        return p
    raise FileNotFoundError(f"Could not find a raw_data file for {number!r} in {RAW_DIR}")


def _zero_crossing_strain(
    time: np.ndarray, stress: np.ndarray, strain: np.ndarray, max_time: float, min_time: float
) -> tuple[float, float] | None:
    """The last point before Stress crosses from positive to negative
    within [max_time, min_time] — closest to zero while still positive.

    `positive` is already in increasing time order. Quantization can make
    several consecutive samples share the exact same minimum positive
    value (a plateau); plain argmin would return the FIRST of those tied
    samples, which can be a couple of samples earlier than the actual
    crossing. Take the LAST index tied for the minimum instead, since
    that's the one immediately before stress goes negative.
    """
    mask = (time >= max_time) & (time <= min_time)
    idxs = np.where(mask)[0]
    positive = idxs[stress[idxs] > 0]
    if len(positive) == 0:
        return None
    min_val = stress[positive].min()
    best = positive[stress[positive] == min_val][-1]
    return float(time[best]), float(strain[best])


def _nearest_value(target_time: float, other_time: np.ndarray, other_value: np.ndarray) -> float:
    idx = np.searchsorted(other_time, target_time)
    candidates = [i for i in (idx - 1, idx) if 0 <= i < len(other_time)]
    best = min(candidates, key=lambda i: abs(other_time[i] - target_time))
    return float(other_value[best])


def _compute_si_zero_crossing(
    raw_path: Path,
) -> tuple[list[tuple[object, float, float, float, float]], str]:
    main_df, ext_df, ext_name = read_combined_data(raw_path)
    if ext_df is None:
        raise ValueError(f"{raw_path.name} has no 4th/5th-column signal to compute Si against")

    time = main_df["Time"].to_numpy()
    stress = main_df["Stress"].to_numpy()
    strain = main_df["Strain"].to_numpy()

    cycles: list[Cycle]
    cycles, _threshold, _merged = cycles_for_signal(time, stress, STRESS_THRESHOLD, SHOW_LAST_CYCLE)

    ext_time = ext_df["Time"].to_numpy()
    ext_value = ext_df[ext_name].to_numpy()

    rows = []
    for cycle in cycles:
        if cycle.max_value <= 0 or cycle.min_value >= 0:
            continue
        crossing = _zero_crossing_strain(time, stress, strain, cycle.max_time, cycle.min_time)
        if crossing is None:
            continue
        t_zero, strain_i = crossing
        ext_at_zero = _nearest_value(t_zero, ext_time, ext_value)
        si = (ext_at_zero / 2) - (strain_i * 70)
        rows.append((cycle.cycle_number, t_zero, strain_i, ext_at_zero, si))
    return rows, ext_name


def _compute_si_lvdt_max(
    raw_path: Path, cycle_label_fn
) -> tuple[list[tuple[object, float, float, float, float]], str]:
    main_df, ext_df, ext_name = read_combined_data(raw_path)
    if ext_df is None:
        raise ValueError(f"{raw_path.name} has no 4th/5th-column signal to compute Si against")

    main_time = main_df["Time"].to_numpy()
    strain = main_df["Strain"].to_numpy()

    ext_time = ext_df["Time"].to_numpy()
    ext_value = ext_df[ext_name].to_numpy()

    ext_cycles: list[Cycle]
    ext_cycles, _threshold, _merged = cycles_for_signal(
        ext_time, ext_value, EXTENSION_THRESHOLD, SHOW_LAST_CYCLE
    )

    rows = []
    for cycle in ext_cycles:
        lvdt_max = cycle.max_value
        lvdt_max_time = cycle.max_time
        strain_i = _nearest_value(lvdt_max_time, main_time, strain)
        si = (lvdt_max / 2) - (strain_i * 70)
        label = cycle_label_fn(cycle.cycle_number) if cycle_label_fn else cycle.cycle_number
        rows.append((label, lvdt_max_time, strain_i, lvdt_max, si))
    return rows, ext_name


def compute_si_rows(number: int, raw_path: Path) -> tuple[list[tuple[object, float, float, float, float]], str]:
    if number in ZERO_CROSSING_FILE_NUMBERS:
        return _compute_si_zero_crossing(raw_path)
    cycle_label_fn = paired_cycle_label if number in PAIRED_CYCLE_FILE_NUMBERS else None
    return _compute_si_lvdt_max(raw_path, cycle_label_fn)


def _write_si_sheet_and_chart(
    output_path: Path, rows: list[tuple[object, float, float, float, float]], ext_name: str
) -> None:
    wb = load_workbook(output_path)
    if "Ci" in wb.sheetnames:
        del wb["Ci"]  # legacy sheet name from before the Ci -> Si rename
    if "Si" in wb.sheetnames:
        del wb["Si"]
    ws = wb.create_sheet("Si")

    headers = ["Cycle", "Time", "Strain", ext_name, "Si"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(list(row))
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    ws.freeze_panes = "A2"

    charts_ws = wb["Charts"]

    def _chart_title(c) -> str | None:
        try:
            return c.title.tx.rich.p[0].r[0].t
        except (AttributeError, IndexError, TypeError):
            return None

    charts_ws._charts = [c for c in charts_ws._charts if _chart_title(c) != "Ci vs Time"]
    next_row = 1 + 25 * len(charts_ws._charts)

    chart = ScatterChart()
    chart.title = "Si vs Time"
    chart.style = 13
    chart.x_axis.title = "Time"
    chart.y_axis.title = "Si"
    chart.width = 24
    chart.height = 12

    last_row = 1 + len(rows)
    xvalues = Reference(ws, min_col=2, min_row=2, max_row=last_row)
    yvalues = Reference(ws, min_col=5, min_row=1, max_row=last_row)
    series = Series(yvalues, xvalues, title_from_data=True)
    series.marker.symbol = "circle"
    series.marker.size = 6
    series.marker.graphicalProperties.solidFill = SI_SERIES_COLOR
    series.marker.graphicalProperties.line.solidFill = SI_SERIES_COLOR
    series.smooth = False
    series.graphicalProperties.line.width = 15000
    series.graphicalProperties.line.solidFill = SI_SERIES_COLOR
    chart.series.append(series)
    chart.legend.position = "b"
    chart.legend.overlay = False

    charts_ws.add_chart(chart, f"A{next_row}")
    wb.save(output_path)


def main() -> int:
    for number in sorted(ZERO_CROSSING_FILE_NUMBERS | LVDT_MAX_FILE_NUMBERS):
        raw_path = _find_raw_path(number)
        output_path = RESULTS_DIR / f"{raw_path.stem}_result.xlsx"
        if not output_path.exists():
            raise SystemExit(f"{output_path} not found — run run.py first.")

        rows, ext_name = compute_si_rows(number, raw_path)
        _write_si_sheet_and_chart(output_path, rows, ext_name)
        print(f"{raw_path.name}: added Si sheet with {len(rows)} rows -> {output_path.name}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
