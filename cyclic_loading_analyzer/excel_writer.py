"""Excel workbook output: Raw Data, one Results sheet per signal, Charts.

Signals are organized into groups that share one Time column (e.g. Stress
and Strain, sampled by the same DAQ). A workbook can have more than one
group with independent time bases and row counts (e.g. a second group for
an LVDT-sampled Extension signal) — each still gets its own pair of
charts, all landing in the same "Charts" sheet.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .detector import Cycle

CHART_MAX_POINTS = 5000
HEADER_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=12)
CENTER = Alignment(horizontal="center", vertical="center")

RAW_SERIES_COLOR = "1F77B4"  # blue
MAX_SERIES_COLOR = "C00000"  # red
MIN_SERIES_COLOR = "FFC000"  # amber (more legible than pure yellow on white)
STRAIN_OVERLAY_COLOR = "C00000"  # red

TIME_AXIS_TITLE = "Time (min)"


def unit_for(name: str) -> str:
    """The physical unit for a signal name, or "" if it's dimensionless.

    Stress is MPa; Strain is a dimensionless ratio, so it gets no unit
    suffix; any other signal in this codebase is the independently-
    sampled third signal (e.g. Extension from an LVDT), measured in mm.
    """
    if name == "Stress":
        return "MPa"
    if name == "Strain":
        return ""
    return "mm"


def axis_label_with_unit(name: str) -> str:
    """A signal's axis/column label with its unit appended, e.g. "Stress (MPa)" -
    or just the bare name for a dimensionless signal like Strain."""
    unit = unit_for(name)
    return f"{name} ({unit})" if unit else name


@dataclass(frozen=True)
class SignalGroup:
    """One or more signals sharing a single Time column.

    `time` is that shared time base. `signals` is a list of
    (name, values, cycles) — `values` is the raw signal aligned to `time`
    (same length), `cycles` are its already-detected Max/Min pairs.
    `cycle_labels`, if given (same length as `time`), is written as a
    "Cycle" column right before the signal columns — see
    `cycle_labeling.cycle_labels`.
    """

    time: Sequence[float]
    signals: Sequence[tuple[str, Sequence[float], Sequence[Cycle]]]
    cycle_labels: Sequence[object] | None = None


def _write_raw_data_sheet(
    wb: Workbook, groups: Sequence[SignalGroup]
) -> tuple[Worksheet, list[tuple[int, list[int], int]]]:
    """Write each group as its own block of columns (Time, signal...),
    left to right, separated by one blank spacer column. Groups may have
    different row counts — shorter ones just leave blank cells below
    their own last row.

    Returns (worksheet, layout) where layout[i] = (time_col, value_cols,
    n_rows) for groups[i], all 1-indexed column numbers.
    """
    ws = wb.create_sheet("Raw Data")

    layout: list[tuple[int, list[int], int]] = []
    header_row: list[str] = []
    col = 1
    for gi, group in enumerate(groups):
        time_col = col
        value_cols = []
        header_row.append(TIME_AXIS_TITLE)
        col += 1
        if group.cycle_labels is not None:
            header_row.append("Cycle")
            col += 1
        for name, _values, _cycles in group.signals:
            header_row.append(axis_label_with_unit(name))
            value_cols.append(col)
            col += 1
        layout.append((time_col, value_cols, len(group.time)))
        if gi < len(groups) - 1:
            header_row.append("")  # spacer column
            col += 1

    ws.append(header_row)
    for cell in ws[1]:
        if cell.value:
            cell.font = HEADER_FONT

    max_rows = max((n for _, _, n in layout), default=0)
    for r in range(max_rows):
        row: list[object] = []
        for gi, group in enumerate(groups):
            _time_col, _value_cols, n_rows = layout[gi]
            if r < n_rows:
                row.append(group.time[r])
                if group.cycle_labels is not None:
                    row.append(group.cycle_labels[r])
                for _name, values, _cycles in group.signals:
                    row.append(values[r])
            else:
                row.append(None)
                if group.cycle_labels is not None:
                    row.append(None)
                row.extend([None] * len(group.signals))
            if gi < len(groups) - 1:
                row.append(None)  # spacer
        ws.append(row)

    for col_idx in range(1, len(header_row) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 14
    ws.freeze_panes = "A2"
    return ws, layout


def _write_results_sheet(
    wb: Workbook, sheet_name: str, signal_name: str, cycles: Sequence[Cycle]
) -> Worksheet:
    ws = wb.create_sheet(sheet_name)
    unit = unit_for(signal_name)
    max_label = f"Max ({unit})" if unit else "Max"
    min_label = f"Min ({unit})" if unit else "Min"

    ws.merge_cells("A1:C1")
    ws["A1"] = "Max"
    ws.merge_cells("D1:F1")
    ws["D1"] = "Min"
    for coord in ("A1", "D1"):
        ws[coord].font = TITLE_FONT
        ws[coord].alignment = CENTER

    headers = ["Cycle", TIME_AXIS_TITLE, max_label, "Cycle", TIME_AXIS_TITLE, min_label]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    for row_offset, cycle in enumerate(cycles, start=3):
        ws.cell(row=row_offset, column=1, value=cycle.cycle_number)
        ws.cell(row=row_offset, column=2, value=cycle.max_time)
        ws.cell(row=row_offset, column=3, value=cycle.max_value)
        ws.cell(row=row_offset, column=4, value=cycle.cycle_number)
        ws.cell(row=row_offset, column=5, value=cycle.min_time)
        ws.cell(row=row_offset, column=6, value=cycle.min_value)

    for col_idx in range(1, 7):
        ws.column_dimensions[get_column_letter(col_idx)].width = 12
    ws.freeze_panes = "A3"
    return ws


def add_dual_axis_stress_strain_chart(
    source_ws: Worksheet, time_col: int, stress_col: int, strain_col: int, last_row: int
) -> ScatterChart:
    """One large chart overlaying Stress and Strain vs Time on the same
    Time axis, Strain on a secondary Y axis -- their scales are wildly
    different (Stress in the hundreds, Strain a small fraction), so
    without a second axis one of the two would look flat. Lets the two
    shapes be compared directly, point for point in time.
    """
    xvalues = Reference(source_ws, min_col=time_col, min_row=2, max_row=last_row)

    stress_chart = ScatterChart()
    stress_chart.title = "Stress & Strain vs Time"
    stress_chart.style = 13
    stress_chart.x_axis.title = TIME_AXIS_TITLE
    stress_chart.y_axis.title = axis_label_with_unit("Stress")
    stress_chart.width = 48
    stress_chart.height = 24

    stress_values = Reference(source_ws, min_col=stress_col, min_row=1, max_row=last_row)
    stress_series = Series(stress_values, xvalues, title_from_data=True)
    stress_series.marker.symbol = "none"
    stress_series.smooth = False
    stress_series.graphicalProperties.line.width = 12000
    stress_series.graphicalProperties.line.solidFill = RAW_SERIES_COLOR
    stress_chart.series.append(stress_series)
    stress_chart.y_axis.crosses = "min"

    strain_chart = ScatterChart()
    strain_chart.y_axis.axId = 200
    strain_chart.y_axis.title = axis_label_with_unit("Strain")
    strain_chart.y_axis.crosses = "max"

    strain_values = Reference(source_ws, min_col=strain_col, min_row=1, max_row=last_row)
    strain_series = Series(strain_values, xvalues, title_from_data=True)
    strain_series.marker.symbol = "none"
    strain_series.smooth = False
    strain_series.graphicalProperties.line.width = 12000
    strain_series.graphicalProperties.line.solidFill = STRAIN_OVERLAY_COLOR
    strain_chart.series.append(strain_series)

    stress_chart += strain_chart
    stress_chart.legend.position = "b"
    stress_chart.legend.overlay = False
    return stress_chart


def _add_charts_sheet(
    wb: Workbook,
    groups: Sequence[SignalGroup],
    layout: list[tuple[int, list[int], int]],
    stress_strain_overlay: bool = False,
) -> Worksheet:
    ws = wb.create_sheet("Charts")
    raw_ws = wb["Raw Data"]
    helper_ws = None  # created lazily, only if some group needs downsampling

    def add_raw_series(chart: ScatterChart, source_ws: Worksheet, x_col: int, y_col: int, last_row: int) -> None:
        xvalues = Reference(source_ws, min_col=x_col, min_row=2, max_row=last_row)
        yvalues = Reference(source_ws, min_col=y_col, min_row=1, max_row=last_row)
        series = Series(yvalues, xvalues, title_from_data=True)
        series.marker.symbol = "none"
        series.smooth = False
        series.graphicalProperties.line.width = 9000
        series.graphicalProperties.line.solidFill = RAW_SERIES_COLOR
        chart.series.append(series)

    def add_max_min_series(chart: ScatterChart, results_ws: Worksheet, n_cycles: int) -> None:
        if n_cycles == 0:
            return
        last_row = 2 + n_cycles
        # Max: Time in col B, value in col C (red). Min: Time in col E, value in col F (amber).
        for x_col, y_col, color in ((2, 3, MAX_SERIES_COLOR), (5, 6, MIN_SERIES_COLOR)):
            xvalues = Reference(results_ws, min_col=x_col, min_row=3, max_row=last_row)
            yvalues = Reference(results_ws, min_col=y_col, min_row=2, max_row=last_row)
            series = Series(yvalues, xvalues, title_from_data=True)
            series.marker.symbol = "circle"
            series.marker.size = 6
            series.marker.graphicalProperties.solidFill = color
            series.marker.graphicalProperties.line.solidFill = color
            series.smooth = False
            series.graphicalProperties.line.width = 15000
            series.graphicalProperties.line.solidFill = color
            chart.series.append(series)

    def make_chart(
        title: str,
        y_label: str,
        *,
        raw_source: tuple[Worksheet, int, int, int] | None,
        results_ws: Worksheet,
        n_cycles: int,
    ) -> ScatterChart:
        chart = ScatterChart()
        chart.title = title
        chart.style = 13
        chart.x_axis.title = TIME_AXIS_TITLE
        chart.y_axis.title = axis_label_with_unit(y_label)
        chart.width = 24
        chart.height = 12
        if raw_source is not None:
            source_ws, x_col, y_col, last_row = raw_source
            add_raw_series(chart, source_ws, x_col, y_col, last_row)
        add_max_min_series(chart, results_ws, n_cycles)
        chart.legend.position = "b"
        chart.legend.overlay = False
        return chart

    combined_charts = []
    maxmin_charts = []
    overlay_chart = None

    for gi, group in enumerate(groups):
        time_col, value_cols, n_rows = layout[gi]
        step = max(1, (n_rows // CHART_MAX_POINTS) + 1) if n_rows > CHART_MAX_POINTS else 1
        last_row = n_rows + 1

        if step == 1:
            source_ws = raw_ws
            source_time_col = time_col
            source_value_cols = value_cols
            source_last_row = last_row
        else:
            # Reference() cannot stride rows, so downsampled chart series
            # are built from a hidden helper sheet with every Nth row of
            # this group only.
            if helper_ws is None:
                helper_ws = wb.create_sheet("_chart_data")
                helper_ws.sheet_state = "hidden"
            base_col = helper_ws.max_column + 1 if helper_ws.max_column > 1 else 1
            names = [name for name, _v, _c in group.signals]
            for i, h in enumerate(["Time"] + names):
                helper_ws.cell(row=1, column=base_col + i, value=h)
            out_row = 2
            for r in range(1, n_rows, step):
                helper_ws.cell(row=out_row, column=base_col, value=raw_ws.cell(row=r + 1, column=time_col).value)
                for i, vcol in enumerate(value_cols):
                    helper_ws.cell(
                        row=out_row, column=base_col + 1 + i, value=raw_ws.cell(row=r + 1, column=vcol).value
                    )
                out_row += 1
            source_ws = helper_ws
            source_time_col = base_col
            source_value_cols = [base_col + 1 + i for i in range(len(value_cols))]
            source_last_row = out_row - 1

        for i, (name, _values, _cycles) in enumerate(group.signals):
            results_ws = wb[f"{name} Results"]
            n_cycles = max(0, results_ws.max_row - 2)
            raw_source = (source_ws, source_time_col, source_value_cols[i], source_last_row)
            combined_charts.append(
                make_chart(
                    f"{name} vs Time (with per-cycle Max & Min)",
                    name,
                    raw_source=raw_source,
                    results_ws=results_ws,
                    n_cycles=n_cycles,
                )
            )
            maxmin_charts.append(
                make_chart(
                    f"{name} Max & Min vs Time",
                    name,
                    raw_source=None,
                    results_ws=results_ws,
                    n_cycles=n_cycles,
                )
            )

        if gi == 0 and stress_strain_overlay:
            # group 0 is always [Stress, Strain] sharing one Time column.
            overlay_chart = add_dual_axis_stress_strain_chart(
                source_ws, source_time_col, source_value_cols[0], source_value_cols[1], source_last_row
            )

    charts = ([overlay_chart] if overlay_chart is not None else []) + combined_charts + maxmin_charts
    row = 1
    for chart in charts:
        ws.add_chart(chart, f"A{row}")
        row += round(25 * chart.height / 12)  # 25 rows of spacing per 12cm of chart height

    return ws


def _write_simple_table_sheet(wb: Workbook, sheet_name: str, rows: Sequence[dict]) -> Worksheet:
    ws = wb.create_sheet(sheet_name)
    if not rows:
        return ws
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in rows:
        ws.append([row[h] for h in headers])
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16
    ws.freeze_panes = "A2"
    return ws


def _write_combined_cross_reference_sheet(wb: Workbook, tables: dict[str, Sequence[dict]]) -> Worksheet:
    """All cross-reference tables laid out side by side in one sheet."""
    ws = wb.create_sheet("Combined")
    col = 1
    for name, rows in tables.items():
        if not rows:
            continue
        headers = list(rows[0].keys())
        title_cell = ws.cell(row=1, column=col, value=name)
        title_cell.font = TITLE_FONT
        for i, h in enumerate(headers):
            cell = ws.cell(row=2, column=col + i, value=h)
            cell.font = HEADER_FONT
        for r, row in enumerate(rows, start=3):
            for i, h in enumerate(headers):
                ws.cell(row=r, column=col + i, value=row[h])
        for i in range(len(headers)):
            ws.column_dimensions[get_column_letter(col + i)].width = 14
        col += len(headers) + 1  # +1 blank spacer column between tables
    ws.freeze_panes = "A3"
    return ws


SI_SERIES_COLOR = "1F77B4"


def _write_si_sheet_and_chart(
    wb: Workbook,
    charts_ws: Worksheet,
    rows: Sequence[tuple[object, float, float, float, float]],
    ext_name: str,
) -> None:
    """Add the "Si" sheet (Cycle, Time, Strain, Extension, Si) and an
    "Si vs Cycle" chart, to an already-open workbook (before it's saved).

    A LineChart with a category axis, not ScatterChart (which needs a
    numeric X): the Cycle column holds text labels ("1_1", "1_2", ...)
    for the paired files, so Si is plotted one point per cycle, evenly
    spaced by cycle order rather than by its numeric/time value.
    """
    ws = wb.create_sheet("Si")
    headers = ["Cycle", TIME_AXIS_TITLE, axis_label_with_unit("Strain"), axis_label_with_unit(ext_name), "Si"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(list(row))
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 15
    ws.freeze_panes = "A2"

    next_row = 1 + 25 * len(charts_ws._charts)

    chart = LineChart()
    chart.title = "Si vs Cycle"
    chart.style = 13
    chart.x_axis.title = "Cycle"
    chart.y_axis.title = "Si"
    chart.width = 24
    chart.height = 12

    last_row = 1 + len(rows)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)
    data = Reference(ws, min_col=5, min_row=1, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    series = chart.series[0]
    series.marker.symbol = "circle"
    series.marker.size = 6
    series.marker.graphicalProperties.solidFill = SI_SERIES_COLOR
    series.marker.graphicalProperties.line.solidFill = SI_SERIES_COLOR
    series.smooth = False
    series.graphicalProperties.line.width = 15000
    series.graphicalProperties.line.solidFill = SI_SERIES_COLOR
    chart.legend.position = "b"
    chart.legend.overlay = False

    charts_ws.add_chart(chart, f"A{next_row}")


def write_workbook(
    output_path: str | Path,
    groups: Sequence[SignalGroup],
    cross_reference_tables: dict[str, Sequence[dict]] | None = None,
    si_rows: Sequence[tuple[object, float, float, float, float]] | None = None,
    si_ext_name: str | None = None,
    stress_strain_overlay: bool = False,
) -> None:
    """Write Raw Data, one Results sheet per signal, and Charts.

    Each group in `groups` shares one Time column; groups may have
    different row counts (independent sampling devices). Two charts are
    produced per signal (raw-with-overlay, and Max/Min-only) — e.g. 4
    charts total for a Stress+Strain group, 2 more for an Extension group.

    `cross_reference_tables`, if given, adds one sheet per table plus a
    "Combined" sheet with all of them side by side — see
    `cross_reference.build_cross_reference_tables`.

    `si_rows`/`si_ext_name`, if given, add the "Si" sheet + chart in this
    same save — see `_write_si_sheet_and_chart` — instead of a caller
    re-opening the file afterward to append it.

    `stress_strain_overlay`, if true, adds one large chart overlaying
    Stress and Strain vs Time on a shared Time axis (Strain on a
    secondary Y axis) as the first chart in "Charts" — see
    `add_dual_axis_stress_strain_chart`.
    """
    wb = Workbook()
    wb.remove(wb.active)  # drop default empty sheet

    _raw_ws, layout = _write_raw_data_sheet(wb, groups)
    for group in groups:
        for name, _values, cycles in group.signals:
            _write_results_sheet(wb, f"{name} Results", name, cycles)
    charts_ws = _add_charts_sheet(wb, groups, layout, stress_strain_overlay)

    if cross_reference_tables:
        for name, rows in cross_reference_tables.items():
            _write_simple_table_sheet(wb, name, rows)
        _write_combined_cross_reference_sheet(wb, cross_reference_tables)

    if si_rows is not None:
        _write_si_sheet_and_chart(wb, charts_ws, si_rows, si_ext_name)

    wb.save(str(output_path))
